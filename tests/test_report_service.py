"""Excel-отчёт мастера: какие записи попадают в период, как считаются доход и счётчики клиентов."""

import datetime as dt
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from app.database.models import Booking, User
from app.services.report_service import REPORT_HEADER, build_report_workbook
from tests.conftest import MASTER_ID, RIGA_TZ, make_client

D = dt.date
WED = D(2030, 1, 16)
LONG_AGO = dt.datetime(2029, 6, 1, 12, 0, tzinfo=dt.timezone.utc)  # создана/изменена вне любого тестового периода


def local(hour: int, minute: int = 0, day: dt.date = WED, tz=RIGA_TZ) -> dt.datetime:
    return dt.datetime.combine(day, dt.time(hour, minute), tzinfo=tz)


async def add_booking(
    session,
    user: User,
    service,
    start: dt.datetime,
    *,
    status: str = "confirmed",
    price="40",
    service_name: str | None = None,
    created_at: dt.datetime = LONG_AGO,
    updated_at: dt.datetime = LONG_AGO,
    manual_name: str | None = None,
    manual_contact: str | None = None,
) -> Booking:
    booking = Booking(master_id=MASTER_ID,
        user_id=user.id,
        service_id=service.id,
        start_datetime=start,
        end_datetime=start + dt.timedelta(hours=1),
        status=status,
        service_name_snapshot=service_name or service.name,
        service_price_snapshot=price,
        service_duration_snapshot=60,
        created_at=created_at,
        updated_at=updated_at,
        manual_client_name=manual_name,
        manual_client_contact=manual_contact,
    )
    session.add(booking)
    await session.commit()
    return booking


async def report(session, period="today", now=None, tz=RIGA_TZ):
    buffer, filename = await build_report_workbook(session, MASTER_ID, tz, period, now or local(12))
    ws = load_workbook(BytesIO(buffer.getvalue())).active
    rows = list(ws.iter_rows(values_only=True))
    return ws, rows[0], rows[1:-2], rows[-2:], filename


def col(rows, name: str) -> list:
    index = REPORT_HEADER.index(name)
    return [r[index] for r in rows]


async def test_empty_period_has_header_and_zero_totals(session, master_profile):
    ws, header, body, totals, filename = await report(session)

    assert ws.title == "Отчёт"
    assert list(header) == REPORT_HEADER
    assert body == []
    assert [(t[0], t[3]) for t in totals] == [("Фактический доход", "0 BYN"), ("Ожидаемый доход", "0 BYN")]
    assert filename == "otchet_16012030-16012030.xlsx"


@pytest.mark.parametrize(
    ("period", "expected"),
    [("week", "otchet_14012030-20012030.xlsx"), ("month", "otchet_01012030-31012030.xlsx")],
)
async def test_filename_reflects_period(session, master_profile, period, expected):
    assert (await report(session, period))[4] == expected


async def test_day_boundaries_follow_the_masters_timezone(session, master_profile, test_user, manicure_service):
    """«Сегодня» — по местному времени мастера, а не по UTC: 00:30 по Риге — это ещё вчера в UTC."""
    for start in (local(23, 30, WED - dt.timedelta(days=1)), local(0, 30), local(23, 30), local(0, 30, WED + dt.timedelta(days=1))):
        await add_booking(session, test_user, manicure_service, start)

    _, _, body, _, _ = await report(session)

    assert col(body, "Время") == ["00:30–01:30", "23:30–00:30"]
    assert col(body, "Дата") == ["16.01.2030", "16.01.2030"]


async def test_day_boundaries_across_dst_change(session, master_profile, test_user, manicure_service):
    """31 марта 2030 Рига переходит на летнее время: сутки длиной 23 часа, границы не должны «поплыть»."""
    riga = ZoneInfo("Europe/Riga")
    day = D(2030, 3, 31)
    for start in (local(23, 30, day - dt.timedelta(days=1), riga), local(0, 30, day, riga), local(23, 0, day, riga), local(0, 10, day + dt.timedelta(days=1), riga)):
        await add_booking(session, test_user, manicure_service, start)

    _, _, body, _, _ = await report(session, now=local(12, day=day, tz=riga), tz=riga)

    assert col(body, "Время") == ["00:30–01:30", "23:00–00:00"]
    assert col(body, "Дата") == ["31.03.2030", "31.03.2030"]


async def test_rows_are_sorted_by_visit_time(session, master_profile, test_user, manicure_service):
    for hour in (15, 9, 12):
        await add_booking(session, test_user, manicure_service, local(hour))
    _, _, body, _, _ = await report(session)
    assert col(body, "Время") == ["09:00–10:00", "12:00–13:00", "15:00–16:00"]


async def test_record_created_in_period_for_another_day_is_listed_but_earns_nothing(
    session, master_profile, test_user, manicure_service
):
    await add_booking(session, test_user, manicure_service, local(11, day=D(2030, 1, 25)), created_at=local(10))
    _, _, body, totals, _ = await report(session)

    assert col(body, "Дата") == ["25.01.2030"]
    assert col(body, "Создана") == ["16.01.2030 10:00"]
    assert [t[3] for t in totals] == ["0 BYN", "0 BYN"]


async def test_cancellation_in_period_is_listed_with_cancel_time_and_counted(
    session, master_profile, test_user, manicure_service
):
    await add_booking(
        session, test_user, manicure_service, local(11, day=D(2030, 1, 25)), status="cancelled", updated_at=local(9, 30)
    )
    # отменена давно и визит не в периоде — в отчёт не попадает
    await add_booking(session, test_user, manicure_service, local(11, day=D(2030, 1, 26)), status="cancelled")

    _, _, body, totals, _ = await report(session)

    assert len(body) == 1
    assert col(body, "Статус") == ["Отменена"]
    assert col(body, "Отменена") == ["16.01.2030 09:30"]
    assert col(body, "Отмен") == [1] and col(body, "Посещений") == [0]
    assert [t[3] for t in totals] == ["0 BYN", "0 BYN"]


async def test_non_cancelled_booking_has_no_cancel_time_even_if_updated_in_period(
    session, master_profile, test_user, manicure_service
):
    await add_booking(session, test_user, manicure_service, local(11), updated_at=local(9))
    _, _, body, _, _ = await report(session)
    assert col(body, "Отменена") in (["", ], [None])


async def test_revenue_counts_only_visits_of_the_period_by_status(session, master_profile, test_user, manicure_service):
    await add_booking(session, test_user, manicure_service, local(10), status="completed", price="40.50")
    await add_booking(session, test_user, manicure_service, local(14), status="confirmed", price="25")
    await add_booking(session, test_user, manicure_service, local(15), status="no_show", price="99")
    await add_booking(session, test_user, manicure_service, local(16), status="cancelled", price="77")
    # создана сегодня, но визит в другой день — в таблице есть, в доходе нет
    await add_booking(session, test_user, manicure_service, local(10, day=D(2030, 1, 30)), status="completed", price="500", created_at=local(8))

    _, _, body, totals, _ = await report(session)

    assert len(body) == 5
    assert [(t[0], t[3]) for t in totals] == [("Фактический доход", "40.50 BYN"), ("Ожидаемый доход", "65.50 BYN")]
    assert "40.50 BYN" in col(body, "Сумма")


async def test_client_counters_are_per_client_and_per_period(
    session, master_profile, test_user, test_user_2, manicure_service
):
    for hour, status in [(9, "confirmed"), (10, "completed"), (11, "completed"), (12, "cancelled"), (13, "no_show")]:
        await add_booking(session, test_user, manicure_service, local(hour), status=status)
    await add_booking(session, test_user_2, manicure_service, local(14))
    # визит прошлого месяца за пределами периода не должен раздувать счётчики
    await add_booking(session, test_user, manicure_service, local(9, day=D(2029, 12, 20)), status="completed")

    _, _, body, _, _ = await report(session)

    by_client = {}
    for row in body:
        by_client.setdefault(row[REPORT_HEADER.index("Клиент")], set()).add(
            (row[REPORT_HEADER.index("Посещений")], row[REPORT_HEADER.index("Отмен")])
        )
    assert by_client == {"Ivan": {(3, 1)}, "Maria": {(1, 0)}}


async def test_manual_clients_are_grouped_by_name_and_contact_not_by_masters_user(
    session, master_profile, test_user, manicure_service
):
    """Записи «вручную» висят на user_id мастера: счётчики должны считаться по manual_client_*, а не слипаться."""
    await add_booking(session, test_user, manicure_service, local(9), manual_name="Аня", manual_contact="+375291111111")
    await add_booking(session, test_user, manicure_service, local(10), manual_name="Аня", manual_contact="+375291111111")
    await add_booking(session, test_user, manicure_service, local(11), manual_name="Аня", manual_contact="+375292222222")
    await add_booking(session, test_user, manicure_service, local(12), manual_name="Оля", status="cancelled")
    await add_booking(session, test_user, manicure_service, local(13))  # обычная запись самого пользователя

    _, _, body, _, _ = await report(session)

    result = {(r[REPORT_HEADER.index("Клиент")], r[REPORT_HEADER.index("Контакт")]): (r[6], r[7]) for r in body}
    assert result == {
        ("Аня", "+375291111111"): (2, 0),
        ("Аня", "+375292222222"): (1, 0),
        ("Оля", "—"): (0, 1),
        ("Ivan", "@client1"): (1, 0),
    }


async def test_client_name_and_contact_fallbacks(session, master_profile, manicure_service):
    nameless = await make_client(session, 1, username=None, first_name=None, last_name=None)
    full = await make_client(session, 2, username=None, first_name="Анна", last_name="Иванова")
    await session.commit()
    await add_booking(session, nameless, manicure_service, local(9))
    await add_booking(session, full, manicure_service, local(10))

    _, _, body, _, _ = await report(session)

    assert col(body, "Клиент") == ["—", "Анна Иванова"]
    assert col(body, "Контакт") == ["—", "—"]


async def test_status_cells_are_highlighted(session, master_profile, test_user, manicure_service):
    for hour, status in [(9, "cancelled"), (10, "completed"), (11, "confirmed"), (12, "no_show")]:
        await add_booking(session, test_user, manicure_service, local(hour), status=status)

    ws, *_ = await report(session)

    fills = [ws.cell(row=r, column=6).fill for r in range(2, 6)]
    assert fills[0].start_color.rgb.endswith("FFC7CE")
    assert fills[1].start_color.rgb.endswith("C6EFCE")
    assert fills[2].fill_type is None and fills[3].fill_type is None


async def test_user_input_is_never_written_as_an_excel_formula(session, master_profile, manicure_service):
    """Имя в Telegram и название услуги — произвольный текст; «=HYPERLINK(...)» не должен стать формулой в Excel мастера."""
    evil = await make_client(
        session, 7, username="evil", first_name='=HYPERLINK("http://evil.example","Нажми")', last_name=None
    )
    await session.commit()
    await add_booking(
        session, evil, manicure_service, local(10), service_name="=1+1", manual_contact="=SUM(A1:A9)",
    )

    ws, *_ = await report(session)

    for cell in ws[2]:
        assert cell.data_type != "f", f"{cell.coordinate}: {cell.value!r} записано как формула"
    assert ws["C2"].value == "=1+1"
    assert ws["E2"].value.startswith("=HYPERLINK")
