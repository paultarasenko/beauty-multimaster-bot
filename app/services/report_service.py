import datetime as dt
from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database.models import Booking, BookingStatus, MasterClient, User
from app.services.stats_service import get_period_dates
from app.utils.master_settings import DEFAULT_CURRENCY
from app.utils.money import format_money
from app.utils.timezone import to_local

STATUS_TITLES_RU = {
    BookingStatus.CONFIRMED.value: "Активна",
    BookingStatus.CANCELLED.value: "Отменена",
    BookingStatus.COMPLETED.value: "Выполнена",
    BookingStatus.NO_SHOW.value: "Не пришёл",
}

# Колонки, где значения должны быть прижаты к левому краю (а не к правому, как обычно для чисел).
_LEFT_ALIGNED_COLUMNS = ("G", "H")

# Подсветка ячейки статуса: отменена — светло-красный, выполнена — светло-зелёный.
# Активная запись и остальные статусы (например "не пришёл") не подсвечиваются.
_STATUS_FILLS = {
    BookingStatus.CANCELLED.value: PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    BookingStatus.COMPLETED.value: PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
}

REPORT_HEADER = [
    "Дата",
    "Время",
    "Услуга",
    "Сумма",
    "Клиент",
    "Статус",
    "Посещений",
    "Отмен",
    "Контакт",
    "Создана",
    "Отменена",
]

# Визиты, которые реально состоялись или ещё предстоят — учитываются в счётчике "Посещений".
COUNTED_VISIT_STATUSES = (BookingStatus.CONFIRMED.value, BookingStatus.COMPLETED.value)


@dataclass
class _ClientAgg:
    visits: int = 0
    cancels: int = 0


def _client_display_name(user: User | None, manual_name: str | None) -> str:
    if manual_name:
        return manual_name
    if user is None:
        return "—"
    full = " ".join(part for part in (user.first_name, user.last_name) if part)
    return full or "—"


def _client_contact(user: User | None, manual_contact: str | None, is_manual: bool = False) -> str:
    if manual_contact:
        return manual_contact
    if is_manual:
        # Запись ручная: user_id — это владелец-мастер, его @username не контакт клиента.
        return "—"
    if user and user.username:
        return f"@{user.username}"
    return "—"


def _client_key(booking: Booking) -> str:
    """Ключ группировки клиента. Записи, добавленные мастером вручную для клиента без своего
    аккаунта в боте, формально висят на user_id мастера — поэтому для них группируем по
    контакту/имени из manual_client_*, а не по user_id (иначе все такие клиенты слились бы
    в одну строку "мастер")."""
    if booking.manual_client_name:
        return f"manual:{booking.manual_client_name}|{booking.manual_client_contact or ''}"
    return f"user:{booking.user_id}"


async def _fetch_bookings(
    session: AsyncSession, master_id: int, range_start: dt.datetime, range_end: dt.datetime
) -> list[Booking]:
    """Вся активность за период: визиты на эти даты, записи, созданные за период (на любую дату),
    и отмены, сделанные за период (updated_at отменённой брони — момент отмены)."""
    result = await session.execute(
        select(Booking)
        .where(
            Booking.master_id == master_id,
            or_(
                and_(Booking.start_datetime >= range_start, Booking.start_datetime < range_end),
                and_(Booking.created_at >= range_start, Booking.created_at < range_end),
                and_(
                    Booking.status == BookingStatus.CANCELLED.value,
                    Booking.updated_at >= range_start,
                    Booking.updated_at < range_end,
                ),
            )
        )
        .order_by(Booking.start_datetime)
    )
    return list(result.scalars().all())


async def _fetch_users_by_ids(session: AsyncSession, master_id: int, user_ids: set[int]) -> dict[int, User]:
    """Клиенты этого мастера по id (чужие клиенты в отчёт не попадают)."""
    if not user_ids:
        return {}
    result = await session.execute(
        select(User)
        .join(MasterClient, MasterClient.user_id == User.id)
        .where(MasterClient.master_id == master_id, User.id.in_(user_ids))
    )
    return {u.id: u for u in result.scalars().all()}


def _force_text_cells(ws) -> None:
    """openpyxl превращает любую строку, начинающуюся с «=», в формулу. В отчёт попадает пользовательский текст
    (имя в Telegram, название услуги, контакт), поэтому «=HYPERLINK(...)» в имени клиента исполнялся бы в Excel
    мастера. Формул в отчёте нет вообще — все такие ячейки принудительно делаем текстовыми."""
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                cell.data_type = "s"


def _autosize_columns(ws, rows: list[list], header: list[str]) -> None:
    widths = [len(str(h)) for h in header]
    for row in rows:
        for i, value in enumerate(row):
            widths[i] = max(widths[i], len(str(value)))
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)


async def build_report_workbook(
    session: AsyncSession, master_id: int, tz: dt.tzinfo, period: str, now: dt.datetime, currency: str = DEFAULT_CURRENCY
) -> tuple[BytesIO, str]:
    """Строит .xlsx с одним листом «Отчёт»: колонки A–K (Дата, Время, Услуга, Сумма, Клиент,
    Статус, Посещений, Отмен, Контакт, Создана, Отменена), по одной строке на бронь, затронутую
    за выбранный period: визит в этот период, создание записи или её отмена в этот период.

    «Посещений» и «Отмен» — счётчики по клиенту в рамках того же period (не за всё время).
    Внизу — две строки: «Фактический доход» (сумма выполненных записей) и
    «Ожидаемый доход» (фактическая плюс активные, ещё не выполненные записи).

    Возвращает (файл в памяти, имя файла).
    """
    start_date, end_date = get_period_dates(period, now.date())
    range_start = dt.datetime.combine(start_date, dt.time.min, tzinfo=tz)
    range_end = dt.datetime.combine(end_date + dt.timedelta(days=1), dt.time.min, tzinfo=tz)

    period_bookings = await _fetch_bookings(session, master_id, range_start, range_end)

    needed_user_ids = {b.user_id for b in period_bookings}
    users_by_id = await _fetch_users_by_ids(session, master_id, needed_user_ids)

    # Посещения/отмены по клиенту — считаем только по броням из этого же периода.
    client_counts: dict[str, _ClientAgg] = {}
    for b in period_bookings:
        key = _client_key(b)
        agg = client_counts.setdefault(key, _ClientAgg())
        if b.status in COUNTED_VISIT_STATUSES:
            agg.visits += 1
        elif b.status == BookingStatus.CANCELLED.value:
            agg.cancels += 1

    rows: list[list] = []
    statuses: list[str] = []
    actual_revenue = Decimal("0")
    expected_revenue = Decimal("0")
    for b in period_bookings:
        user = users_by_id.get(b.user_id)
        start_local = to_local(b.start_datetime, tz)
        end_local = to_local(b.end_datetime, tz)
        key = _client_key(b)
        agg = client_counts[key]
        price = Decimal(str(b.service_price_snapshot))
        # Доход — только по визитам выбранного периода (в таблице остаются и записи, созданные/отменённые
        # за период). Фактическая — выполненные; ожидаемая — выполненные + активные (ещё не выполненные).
        if range_start <= b.start_datetime < range_end:
            if b.status == BookingStatus.COMPLETED.value:
                actual_revenue += price
                expected_revenue += price
            elif b.status == BookingStatus.CONFIRMED.value:
                expected_revenue += price
        statuses.append(b.status)
        rows.append(
            [
                start_local.strftime("%d.%m.%Y"),
                f"{start_local.strftime('%H:%M')}–{end_local.strftime('%H:%M')}",
                b.service_name_snapshot,
                format_money(price, currency),
                _client_display_name(user, b.manual_client_name),
                STATUS_TITLES_RU.get(b.status, b.status),
                agg.visits,
                agg.cancels,
                _client_contact(user, b.manual_client_contact, is_manual=bool(b.manual_client_name)),
                to_local(b.created_at, tz).strftime("%d.%m.%Y %H:%M"),
                (
                    to_local(b.updated_at, tz).strftime("%d.%m.%Y %H:%M")
                    if b.status == BookingStatus.CANCELLED.value
                    else ""
                ),
            ]
        )

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    ws.sheet_view.showRowColHeaders = False
    ws.freeze_panes = "A2"
    ws.append(REPORT_HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row_idx, row in enumerate(rows, start=2):
        ws.append(row)
        fill = _STATUS_FILLS.get(statuses[row_idx - 2])
        if fill is not None:
            ws.cell(row=row_idx, column=6).fill = fill

    total_rows = [
        ["Фактический доход", "", "", format_money(actual_revenue, currency), "", "", "", "", "", "", ""],
        ["Ожидаемый доход", "", "", format_money(expected_revenue, currency), "", "", "", "", "", "", ""],
    ]
    for total_row in total_rows:
        ws.append(total_row)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)

    _force_text_cells(ws)

    left_align = Alignment(horizontal="left")
    for col in _LEFT_ALIGNED_COLUMNS:
        for row_idx in range(2, ws.max_row + 1):
            ws[f"{col}{row_idx}"].alignment = left_align

    _autosize_columns(ws, rows + total_rows, REPORT_HEADER)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"otchet_{start_date.strftime('%d%m%Y')}-{end_date.strftime('%d%m%Y')}.xlsx"
    return buffer, filename